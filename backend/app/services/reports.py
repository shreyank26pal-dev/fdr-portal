"""Management reports, exportable as JSON, Excel (.xlsx) or PDF."""
import io
import re
from datetime import timedelta
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.enums import HELD_STATUSES, PIPELINE_STATUSES, REFUND_OPEN_STATUSES, RefundStatus
from app.models import AuditLog, Institution, Instrument, RefundRequest
from app.services import analytics, clock, rules


class Report:
    def __init__(self, title: str, columns: list[str], rows: list[list[Any]], note: str | None = None):
        self.title, self.columns, self.rows, self.note = title, columns, rows, note

    def as_json(self) -> dict[str, Any]:
        return {"title": self.title, "generated_at": clock.now(), "note": self.note, "columns": self.columns,
                "rows": [dict(zip(self.columns, r)) for r in self.rows]}

    def as_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = re.sub(r"[\\/*?:\[\]]", "-", self.title)[:31]
        ws.append([self.title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Generated {clock.now():%d-%m-%Y %H:%M} UTC" + (f" - {self.note}" if self.note else "")])
        ws.append([])
        ws.append(self.columns)
        for c in range(1, len(self.columns) + 1):
            cell = ws.cell(row=4, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
            cell.alignment = Alignment(wrap_text=True)
        for r in self.rows:
            ws.append([_cell(v) for v in r])
        for i, col in enumerate(self.columns, start=1):
            width = max([len(str(col))] + [len(str(_cell(r[i - 1]))) for r in self.rows[:200]]) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
        ws.freeze_panes = "A5"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def as_pdf(self) -> bytes:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                                topMargin=12 * mm, bottomMargin=12 * mm)
        ss = getSampleStyleSheet()
        small = ss["BodyText"].clone("small", fontSize=7, leading=8)
        data = [[Paragraph(f"<b>{c}</b>", small) for c in self.columns]]
        data += [[Paragraph(str(_cell(v)), small) for v in r] for r in self.rows] or [["No records"] + [""] * (len(self.columns) - 1)]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story = [Paragraph(f"<b>AICTE SecureDeposit - {self.title}</b>", ss["Title"]),
                 Paragraph(f"Generated {clock.now():%d-%m-%Y %H:%M} UTC. {self.note or ''}", ss["Normal"]),
                 Spacer(1, 4 * mm), t]
        doc.build(story)
        return buf.getvalue()


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if hasattr(v, "value"):
        return v.value
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    if hasattr(v, "quantize"):
        return float(v)
    return v


def _scope(q, inst_ids):
    return q.where(Instrument.institution_id.in_(inst_ids)) if inst_ids is not None else q


def maturity_schedule(db: Session, days: int = 90, inst_ids=None, **_) -> Report:
    today = clock.today()
    q = _scope(select(Instrument).where(Instrument.status.in_(HELD_STATUSES),
                                        Instrument.maturity_date.between(today, today + timedelta(days=days)))
               .order_by(Instrument.maturity_date), inst_ids)
    rows = [[i.maturity_date, (i.maturity_date - today).days, i.institution.aicte_id, i.institution.name,
             i.institution.region, i.fdr_number, i.bank_name, i.amount, i.status, i.trust_level]
            for i in db.execute(q).scalars()]
    return Report(f"Maturity schedule - next {days} days",
                  ["Maturity date", "Days left", "AICTE ID", "Institution", "Region", "FDR no", "Bank", "Amount",
                   "Status", "Trust"], rows)


def pending_refunds(db: Session, inst_ids=None, **_) -> Report:
    q = select(RefundRequest).where(RefundRequest.status.in_(REFUND_OPEN_STATUSES)).order_by(RefundRequest.created_at)
    if inst_ids is not None:
        q = q.where(RefundRequest.institution_id.in_(inst_ids))
    now = clock.now()
    rows = []
    for r in db.execute(q).scalars():
        inst = r.instrument
        failing = [c["label"] for c in (r.precondition_results or []) if not c["passed"] and c["blocking"]]
        rows.append([r.ref_number, r.created_at, (now - r.created_at).days, inst.institution.name, inst.fdr_number,
                     inst.amount, r.ground, r.status, "; ".join(failing) or "-"])
    return Report("Pending release / refund requests",
                  ["Ref", "Requested", "Age (days)", "Institution", "FDR no", "Amount", "Ground", "Status",
                   "Failing pre-conditions"], rows)


def release_register(db: Session, inst_ids=None, **_) -> Report:
    q = select(RefundRequest).where(RefundRequest.status == RefundStatus.PROCESSED).order_by(RefundRequest.processed_at)
    if inst_ids is not None:
        q = q.where(RefundRequest.institution_id.in_(inst_ids))
    rows = [[r.ref_number, r.processed_at, r.instrument.institution.name, r.instrument.fdr_number,
             r.instrument.bank_name, r.instrument.amount, r.ground] for r in db.execute(q).scalars()]
    return Report("Release register", ["Ref", "Processed on", "Institution", "FDR no", "Bank", "Amount", "Ground"], rows)


def bank_summary(db: Session, inst_ids=None, **_) -> Report:
    rows = [[r["bank"], r["count"], r["amount"]] for r in analytics.by_bank(db, inst_ids)]
    return Report("Bank-wise exposure (held instruments)", ["Bank", "Instruments", "Amount (INR)"], rows)


def region_summary(db: Session, inst_ids=None, **_) -> Report:
    rows = [[r["region"], r["count"], r["amount"]] for r in analytics.by_region(db, inst_ids)]
    return Report("Region-wise exposure (held instruments)", ["Region", "Instruments", "Amount (INR)"], rows)


def pending_verification(db: Session, inst_ids=None, **_) -> Report:
    now = clock.now()
    q = _scope(select(Instrument).where(Instrument.status.in_(PIPELINE_STATUSES)).order_by(Instrument.submitted_at), inst_ids)
    rows = [[i.ack_number, i.submitted_at, (now - i.submitted_at).days if i.submitted_at else "", i.institution.name,
             i.fdr_number, i.bank_name, i.amount, i.status, i.verification_status, i.risk_score]
            for i in db.execute(q).scalars()]
    return Report("Pending verification - ageing",
                  ["Ack no", "Submitted", "Age (days)", "Institution", "FDR no", "Bank", "Amount", "Status",
                   "Bank verification", "Risk"], rows)


def missing_documents(db: Session, inst_ids=None, **_) -> Report:
    q = _scope(select(Instrument).where(Instrument.status.not_in(["RENEWED", "RELEASED", "REJECTED"]))
               .order_by(Instrument.institution_id), inst_ids)
    rows = []
    for i in db.execute(q).scalars():
        miss = rules.missing_documents(db, i)
        if miss:
            rows.append([i.institution.aicte_id, i.institution.name, i.ack_number or f"draft #{i.id}", i.fdr_number,
                         i.status, ", ".join(miss)])
    return Report("Missing documents", ["AICTE ID", "Institution", "Ack / ID", "FDR no", "Status", "Missing"], rows)


def audit_extract(db: Session, institution_id: int | None = None, inst_ids=None, **_) -> Report:
    q = select(AuditLog).order_by(AuditLog.id)
    if institution_id:
        q = q.where(AuditLog.institution_id == institution_id)
    elif inst_ids is not None:
        q = q.where(AuditLog.institution_id.in_(inst_ids))
    rows = [[a.id, a.ts.strftime("%d-%m-%Y %H:%M:%S"), a.actor_email, a.actor_role, a.action, a.entity_type,
             a.entity_id, a.hash[:16] + "...", a.prev_hash[:16] + "..."] for a in db.execute(q.limit(5000)).scalars()]
    name = db.get(Institution, institution_id).name if institution_id else "all institutions"
    return Report(f"Audit extract - {name}",
                  ["#", "Timestamp", "Actor", "Role", "Action", "Entity", "Entity ID", "Hash", "Prev hash"], rows,
                  note="Hash-chained; verify integrity at /api/v1/audit/verify")


REPORTS: dict[str, tuple[str, Callable[..., Report]]] = {
    "maturity-schedule": ("Upcoming maturities (param: days=30/60/90)", maturity_schedule),
    "pending-refunds": ("Open release / refund requests with failing pre-conditions", pending_refunds),
    "release-register": ("Processed releases", release_register),
    "bank-summary": ("Bank-wise exposure", bank_summary),
    "region-summary": ("Region-wise exposure", region_summary),
    "pending-verification": ("Verification ageing", pending_verification),
    "missing-documents": ("Instruments with missing required documents", missing_documents),
    "audit-extract": ("Audit trail extract (param: institution_id)", audit_extract),
}
